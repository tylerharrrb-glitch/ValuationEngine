"""Reads labelled values from a recalculated workbook: python read_values.py <file.xlsx> <requests.json> -> JSON list."""
import json
import sys

import openpyxl

wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
reqs = json.load(open(sys.argv[2], encoding="utf-8"))
out = []
for r in reqs:
    ws = wb[r["sheet"]]
    val = None
    for row in ws.iter_rows(min_col=2, max_col=2):
        if row[0].value == r["label"]:
            val = ws.cell(row=row[0].row, column=4).value
            break
    out.append(val)
print(json.dumps(out))
