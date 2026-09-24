"""
Excel gate checker (spec 7.3).

  python scripts/excel/check_workbook.py <exported.xlsx> <recalculated.xlsx> <expected.json> [--json out.json]

1. Parity: every expected engine value is found in the recalculated workbook within tolerance
   (0.01 per share, 1 EGP amounts, 1e-7 rates).
2. Checks sheet: every check TRUE; master check TRUE.
3. No error values (#REF!, #DIV/0!, #NAME?, #VALUE!, #N/A, Err:) anywhere.
4. Style lint on the exported file: fonts, colours, fills, merged cells, freeze panes,
   gridlines, hardcoded numbers outside Inputs/Sources/Cover, literal constants in formulas,
   print setup, named ranges, conditional formats.
Prints a full report; exit code 1 on any failure.
"""
import json
import re
import sys

import openpyxl

TOL = {"pershare": 0.01, "amount": 1.0, "rate": 1e-7}
ALLOWED_HARDCODE_SHEETS = {"Inputs", "Sources", "Cover"}
ERROR_TOKENS = ("#REF!", "#DIV/0!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!", "Err:")
REQUIRED_NAMES = ["WACC", "Ke", "Tax_Rate", "Terminal_Growth", "Exit_Multiple", "Shares_Diluted", "Valuation_Date", "Midyear"]


def find_label(ws, label, section=None):
    started = section is None
    for row in ws.iter_rows(min_col=2, max_col=2):
        c = row[0]
        if not started:
            if c.value == section:
                started = True
            continue
        if c.value == label:
            return c.row
    return None


def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def parity(recalc, expected, report):
    failures = 0
    rows = []
    for e in expected["exp"]:
        ws = recalc[e["sheet"]]
        r = find_label(ws, e["label"], e.get("section"))
        got = ws.cell(row=r, column=4 + e.get("col", 0)).value if r else None
        tol = TOL[e["kind"]]
        ok = num(got) and abs(got - e["value"]) <= tol
        failures += 0 if ok else 1
        rows.append((f'{e["sheet"]} / {e["label"]}' + (f' [col {e["col"]}]' if "col" in e else ""), e["value"], got, ok))
    ws = recalc["Sensitivity"]
    for g in expected["grids"]:
        r0 = find_label(ws, g["title"])
        for i in range(5):
            for j in range(5):
                want = g["values"][i][j]
                got = ws.cell(row=r0 + 2 + i, column=4 + j).value if r0 else None
                if want is None:
                    ok = isinstance(got, str)
                else:
                    ok = num(got) and abs(got - want) <= TOL["pershare"]
                failures += 0 if ok else 1
                rows.append((f'Sensitivity / {g["title"]} ({i + 1},{j + 1})', want, got, ok))
    report.append("== Parity: recalculated workbook vs engine ==")
    for name, want, got, ok in rows:
        w = f"{want:,.6f}" if num(want) else str(want)
        gg = f"{got:,.6f}" if num(got) else str(got)
        report.append(f"{'ok ' if ok else 'OUT'}  {name}: engine {w} | excel {gg}")
    report.append(f"parity items: {len(rows)}, outside tolerance: {failures}")
    return failures, len(rows)


def checks_sheet(recalc, report):
    ws = recalc["Checks"]
    start = find_label(ws, "Checks")
    bad = []
    total = 0
    master = None
    for r in range(start + 2, ws.max_row + 1):
        label = ws.cell(row=r, column=2).value
        val = ws.cell(row=r, column=4).value
        if label is None:
            continue
        total += 1
        if label == "Master check":
            master = val
        if val is not True:
            bad.append(f"{label}: {val!r}")
    report.append("== Checks sheet ==")
    report.append(f"checks: {total}, not TRUE: {len(bad)}, master: {master!r}")
    report.extend("  FALSE  " + b for b in bad)
    cover = recalc["Cover"]
    r = find_label(cover, "Master check (all checks TRUE)")
    cover_master = cover.cell(row=r, column=4).value if r else None
    report.append(f"Cover master check: {cover_master!r}")
    return len(bad) + (0 if master is True and cover_master is True else 1), total


def errors(recalc, report):
    found = []
    for ws in recalc.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and any(c.value.startswith(t) for t in ERROR_TOKENS):
                    found.append(f"{ws.title}!{c.coordinate}: {c.value}")
    report.append("== Error values ==")
    report.append(f"error cells: {len(found)}")
    report.extend("  " + f for f in found[:50])
    return len(found)


FUNC = re.compile(r"\b[A-Z][A-Z0-9.]*\s*\(")
STR = re.compile(r'"[^"]*"')
SHEETREF = re.compile(r"('[^']+'|[A-Za-z_][A-Za-z0-9_ ]*)!")
CELL = re.compile(r"\$?[A-Z]{1,3}\$?\d+")
NAME = re.compile(r"\b[A-Za-z_][A-Za-z0-9_]*\b")
NUMBER = re.compile(r"(?<![A-Za-z_$])\d+(?:\.\d+)?(?:[Ee][+-]?\d+)?")


def literals(formula):
    s = STR.sub('""', formula)
    s = SHEETREF.sub("", s)
    s = CELL.sub("R", s)
    s = FUNC.sub("(", s)
    s = NAME.sub("N", s)
    return [m for m in NUMBER.findall(s) if m not in ("0", "1")]


def color_of(font):
    if font is None or font.color is None or font.color.rgb is None:
        return "000000"
    rgb = font.color.rgb
    return rgb[-6:].upper() if isinstance(rgb, str) else "THEME"


def lint(original, report):
    issues = {k: [] for k in ("font", "colour", "fill", "merged", "freeze", "gridlines", "hardcoded", "literal", "print", "names", "condfmt")}
    for ws in original.worksheets:
        if ws.merged_cells.ranges:
            issues["merged"].append(f"{ws.title}: {[str(r) for r in ws.merged_cells.ranges]}")
        if ws.freeze_panes:
            issues["freeze"].append(f"{ws.title}: {ws.freeze_panes}")
        if ws.sheet_view.showGridLines is not False:
            issues["gridlines"].append(ws.title)
        ps = ws.page_setup
        if ps.orientation != "landscape" or ws.sheet_properties.pageSetUpPr is None or not ws.sheet_properties.pageSetUpPr.fitToPage or ps.fitToWidth not in (1, "1"):
            issues["print"].append(f"{ws.title}: orientation={ps.orientation} fitToWidth={ps.fitToWidth}")
        footer = ws.oddFooter.center.text if ws.oddFooter and ws.oddFooter.center else ""
        if not footer or "&P" not in footer or "&N" not in footer or ws.title.replace("&", "&&") not in footer:
            issues["print"].append(f"{ws.title}: footer={footer!r}")
        cf = list(ws.conditional_formatting)
        if cf and ws.title != "Checks":
            issues["condfmt"].append(f"{ws.title}: conditional formatting present")
        for rng in cf:
            for rule in rng.rules:
                if rule.type in ("colorScale", "dataBar", "iconSet"):
                    issues["condfmt"].append(f"{ws.title}: {rule.type}")
                dx = rule.dxf
                if dx is not None and dx.font is not None and color_of(dx.font) != "FF0000":
                    issues["condfmt"].append(f"{ws.title}: conditional font colour {color_of(dx.font)}")
                if dx is not None and dx.fill is not None and dx.fill.fill_type not in (None, "none"):
                    issues["condfmt"].append(f"{ws.title}: conditional fill")
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                if c.font is None or c.font.name != "Calibri":
                    issues["font"].append(f"{ws.title}!{c.coordinate}: {c.font.name if c.font else None}")
                col = color_of(c.font)
                if col not in ("000000", "0000FF"):
                    issues["colour"].append(f"{ws.title}!{c.coordinate}: {col}")
                if c.fill is not None and c.fill.fill_type not in (None, "none"):
                    issues["fill"].append(f"{ws.title}!{c.coordinate}: {c.fill.fill_type}")
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    lits = literals(v[1:])
                    if lits:
                        issues["literal"].append(f"{ws.title}!{c.coordinate}: {lits} in {v[:120]}")
                elif num(v) and ws.title not in ALLOWED_HARDCODE_SHEETS:
                    issues["hardcoded"].append(f"{ws.title}!{c.coordinate}: {v}")
    defined = set(original.defined_names.keys()) if hasattr(original.defined_names, "keys") else {d.name for d in original.defined_names.definedName}
    for n in REQUIRED_NAMES:
        if n not in defined:
            issues["names"].append(f"missing named range {n}")
    report.append("== Style lint (exported workbook) ==")
    labels = {
        "font": "cells not in Calibri", "colour": "font colours outside {000000, 0000FF}", "fill": "cell fills",
        "merged": "merged cells", "freeze": "freeze panes", "gridlines": "sheets with gridlines shown",
        "hardcoded": "hardcoded numbers outside Inputs/Sources/Cover", "literal": "literal constants in formulas (0 and 1 allowed)",
        "print": "print setup deviations", "names": "missing named ranges", "condfmt": "conditional-format deviations",
    }
    total = 0
    for k, items in issues.items():
        report.append(f"{labels[k]}: {len(items)}")
        report.extend("  " + i for i in items[:40])
        if len(items) > 40:
            report.append(f"  ... {len(items) - 40} more")
        total += len(items)
    report.append(f"named ranges defined: {sorted(defined)}")
    report.append(f"sheets: {[ws.title for ws in original.worksheets]}")
    return total, issues


def main():
    orig_path, recalc_path, exp_path = sys.argv[1:4]
    out_json = sys.argv[sys.argv.index("--json") + 1] if "--json" in sys.argv else None
    original = openpyxl.load_workbook(orig_path)
    recalc = openpyxl.load_workbook(recalc_path, data_only=True)
    expected = json.load(open(exp_path, encoding="utf-8"))
    report = []
    pf, pn = parity(recalc, expected, report)
    cf, cn = checks_sheet(recalc, report)
    ef = errors(recalc, report)
    lf, issues = lint(original, report)
    print("\n".join(report))
    summary = {"parityFailures": pf, "parityItems": pn, "checkFailures": cf, "checks": cn, "errorCells": ef,
               "lintIssues": lf, "lint": {k: len(v) for k, v in issues.items()}}
    print(json.dumps(summary))
    if out_json:
        json.dump(summary, open(out_json, "w"), indent=2)
    sys.exit(1 if (pf or cf or ef or lf) else 0)


if __name__ == "__main__":
    main()
